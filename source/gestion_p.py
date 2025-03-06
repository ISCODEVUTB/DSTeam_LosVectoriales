import openpyxl
import os
op_n = "Opcion invalida"
class Paquetes:
    def __init__(self, nombre, peso, precio, tipo, contenido=None, categoria=None, dimension=None, estado_pedido=None, direccion = None, domiciliario = None):
        self.__nombre = nombre
        self.__peso = f"{peso} kg"
        self.__precio = f"{precio} $"
        self.__tipo = tipo
        self.__contenido = contenido if contenido else []
        self.__categoria = categoria if categoria else []
        self.__dimension = dimension if dimension else []
        self.__estado_pedido = estado_pedido
        self.__direccion = direccion
        self.__domiciliario = domiciliario


    @property
    def nombre(self):
        return self.__nombre
    
    @nombre.setter
    def nombre(self, new_name):
        self.__nombre = new_name

    @property
    def peso(self):
        return self.__peso
    
    @peso.setter
    def peso(self, new_peso):
        self.__peso = f"{new_peso} kg"

    @property
    def tipo(self):
        return self.__tipo

    @property
    def precio(self):
        return self.__precio

    @precio.setter
    def precio(self, new_price):
        self.__precio = f"{new_price} $"
    
    @tipo.setter
    def tipo(self, new_tipo):
        self.__tipo = new_tipo

    @property
    def contenido(self):
        return self.__contenido
    
    @contenido.setter
    def contenido(self, cont):
        self.__contenido.pop()
        self.__contenido.append(cont)

    @property
    def categoria(self):
        return self.__categoria
    
    @categoria.setter
    def categoria(self, cat):
        self.__categoria.pop()
        self.__categoria.append(cat)

    @property
    def dimension(self):
        return self.__dimension
    
    @dimension.setter
    def dimension(self, dim):
        self.__dimension.pop()
        self.__dimension.append(dim)

    @property
    def estado_pedido(self):
        return self.__estado_pedido
    
    def agregar_estado_pedido(self, nuevo_estado):
        self.__estado_pedido = nuevo_estado

    @property
    def direccion(self):
        return self.__direccion

    @direccion.setter
    def direccion(self, new_dir):
        self.__direccion = new_dir

    @property
    def domiciliario(self):
        return self.__domiciliario

    @domiciliario.setter
    def domiciliario(self, new_dom):
        self.__domiciliario = new_dom

    def __str__(self):
        return (f"Paquete: {self.__nombre}\n"
                f"precio: {self.__precio}\n"
                f"Peso: {self.__peso}\n"
                f"Tipo: {self.__tipo}\n"
                f"Contenido: {', '.join(self.__contenido) if self.__contenido else 'N/A'}\n"
                f"Categoría: {', '.join(self.__categoria) if self.__categoria else 'N/A'}\n"
                f"Dimensiones: {', '.join(self.__dimension) if self.__dimension else 'N/A'}\n"
                f"estado del pedido: {self.__estado_pedido}\n")
                
class Creacion:
    @staticmethod
    def crear_paquete():

        nombre = input("Nombre del paquete: ")
        precio= input("Digite el precio del paquete: ")
        peso = input("Peso (kg): ")
        tipo = None
        while tipo != "basico" or tipo != "estandar" or tipo != "dimensionado":
            print("El tipo debe ser basico, estandar o dimensionado") 
            tipo = input("Tipo (basico, estandar o dimensionado): ")
            if tipo == "basico" or tipo == "estandar" or tipo == "dimensionado":
                break
        contenido = input("Contenido (separado por comas): ").split(",") 
        categoria = input("Categorías (separadas por comas): ").split(",") 
        dimension = input("Dimensiones (ej: 10x20x30 cm): ").split(",")
        estado_pedido = "pendiente"
        direccion= None
        domiciliario= None
        return Paquetes(nombre, peso, precio, tipo, contenido, categoria, dimension, estado_pedido, direccion, domiciliario)

class ExcelPaquetes:
    FILE_PATH = "paquetes.xlsx"

    @classmethod
    def iniciar_excel(cls):
        if not os.path.exists(cls.FILE_PATH):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append([ "Nombre", "Precio", "Peso", "Tipo", "Contenido", "Categoría", "Dimensiones", "Estado pedido","Direccion","Domiciliario","Usuario"])
            wb.save(cls.FILE_PATH)

    @classmethod
    def guardar(cls, paquete):
        if not os.path.exists(cls.FILE_PATH):
            cls.iniciar_excel()

        wb = openpyxl.load_workbook(cls.FILE_PATH)
        ws = wb.active
        ws.append([
            
            paquete.nombre,
            paquete.precio,
            paquete.peso,
            paquete.tipo,
            ", ".join(paquete.contenido) if paquete.contenido else "N/A",
            ", ".join(paquete.categoria) if paquete.categoria else "N/A",
            ", ".join(paquete.dimension) if paquete.dimension else "N/A",
            paquete.estado_pedido,
            paquete.direccion,
            paquete.domiciliario,
            ])
        
        wb.save(cls.FILE_PATH)
        wb.close()
        print("\nPaquete guardado con éxito en la base de datos.")

    @classmethod
    def mostrar(cls):
        if not os.path.exists(cls.FILE_PATH):
            print("No hay paquetes guardados.")
            return

        wb = openpyxl.load_workbook(cls.FILE_PATH)
        ws = wb.active
        
        if ws.max_row == 1:
            print("No hay paquetes almacenados.")
            return

        print("\nPaquetes almacenados:")
        for row in ws.iter_rows(min_row=2, values_only=True):
            print(f"{row[0]} | {row[1]} | {row[2]} | {row[3]} | {row[4]} | {row[5]} | {row[6]} | {row[7]} | {row[8]}" )
         
        wb.close()
    @classmethod    
    def mostrar_dis(cls):
        if not os.path.exists(cls.FILE_PATH):
            print("No hay paquetes guardados.")
            return

        wb = openpyxl.load_workbook(cls.FILE_PATH)
        ws = wb.active
        
        if ws.max_row == 1:
            print("No hay paquetes almacenados.")
            return

        print("\nPaquetes almacenados:")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[8] == "pendiente":
                print(f"{row[0]} | {row[1]} | {row[2]} | {row[3]} | {row[4]} | {row[5]} | {row[6]} | {row[7]} | Disponible ")
         
        wb.close()


if __name__== "main":
    paquete = Creacion.crear_paquete()  
    ExcelPaquetes.guardar(paquete) 
